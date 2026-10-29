"""Extract first-release dates from the owner's ALFRED downloads (audit report
docs/audit/2026-10-27/12_release_calendar.md).

The owner downloaded "Observations by Real-Time Period" for CPIAUCSL, PAYEMS, PCEPI and ICSA from
alfred.stlouisfed.org on 2026-10-28 (12:38-12:41 PM CDT) and uploaded the four .xlsx files. The raw files are not
committed (FRED terms of use); this script writes only the derived dates: for every observation, the first
realtime_start (the vintage that first carried it) and whether that first vintage had a value.

    python -I docs/audit/2026-10-27/calendar_batch5/extract_alfred.py CPIAUCSL_1.xlsx PAYEMS_1.xlsx PCEPI_1.xlsx ICSA_1.xlsx

Run it with -I: the uploads are untrusted input.
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
OUT = HERE / "sources"
SHEET = "Obs. By Real-Time Period"


def _readme(ws) -> tuple[str, str, list[str]]:
    rows = [r for r in ws.iter_rows(values_only=True)]
    sid = rows[0][0].split(":", 1)[1].strip()
    created = next(r[0] for r in rows if isinstance(r[0], str) and r[0].startswith("File Created:"))
    vintages = [r[0] for r in rows if isinstance(r[0], str) and len(r[0]) == 10 and r[0][4] == "-" and r[1] is None]
    return sid, created.split(":", 1)[1].strip(), vintages


def extract(path: Path) -> tuple[str, str, list[dict], list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sid, created, vintages = _readme(wb["README"])
    first: dict[str, dict] = {}
    for per, value, rt_start, _rt_end in list(wb[SHEET].iter_rows(values_only=True))[1:]:
        if per is None:
            continue
        p, rs = per.date().isoformat(), rt_start.date().isoformat()
        cur = first.get(p)
        if cur is None or rs < cur["first_realtime_start"]:
            first[p] = {"series_id": sid, "period": p, "first_realtime_start": rs, "has_value": value is not None}
        first[p]["vintages"] = first[p].get("vintages", 0) + 1
    return sid, created, [first[p] for p in sorted(first)], vintages


def main(argv: list[str]) -> int:
    OUT.mkdir(exist_ok=True)
    obs, vint = [], []
    for arg in argv:
        sid, created, rows, vintages = extract(Path(arg))
        obs += rows
        vint += [{"series_id": sid, "vintage_date": v, "file_created": created} for v in vintages]
        print(f"{sid}: {len(rows)} observations {rows[0]['period']}..{rows[-1]['period']}, "
              f"{len(vintages)} vintage dates {vintages[0]}..{vintages[-1]} (file created {created})")
    with (OUT / "alfred_first_release.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, ["series_id", "period", "first_realtime_start", "has_value", "vintages"])
        w.writeheader()
        w.writerows(obs)
    with (OUT / "alfred_vintages.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, ["series_id", "vintage_date", "file_created"])
        w.writeheader()
        w.writerows(vint)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
