"""Export the versioned 20-document set and its forms; never export user stores.

Python >=3.11 and python-docx==1.2.0 are required for editable Word copies.
Run from a checkout: python -m tools.compliance.export --output /tmp/qg-delivery
The output directory must be new. It is deliberately outside the source baseline.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
from pathlib import Path
import re
import shutil
import zipfile

from .check import ROOT, check, collect, markdown as report_markdown, registry, safe_path


SUPPLEMENTS = ("README.md", "FACTS.md", "RELEASE_GAPS.md", "SOURCES.md", "PIPELINE.md",
               "registry.json", "release-profile.json", "source-review.json")
PIPELINE_FILES = ("tools/__init__.py", "tools/compliance/__init__.py",
                  "tools/compliance/check.py", "tools/compliance/export.py", "tools/compliance/verify.py",
                  "tests/compliance_acceptance.py", "tests/compliance_export_acceptance.py",
                  ".github/workflows/compliance-review.yml")


def form_section(text: str) -> str:
    match = re.search(r"^## 양식\s*$", text, re.M)
    if match is None:
        raise ValueError("Document has no form section")
    return re.split(r"^## ", text[match.end():], maxsplit=1, flags=re.M)[0].strip()


def inline(text: str, repo_ref: str) -> str:
    """Render the small Markdown vocabulary used by these documents, safely."""
    value = html.escape(text)
    def link(match):
        label, target = match.groups()
        if target.startswith(("https://", "http://")):
            url = target
        elif re.fullmatch(r"\d\d-[a-z-]+\.md", target):
            url = "#doc-" + target[:2]
        elif target in SUPPLEMENTS:
            url = "#" + target.replace(".md", "").lower()
        elif target.startswith("../../") and ".." not in target[6:]:
            url = "https://github.com/xkg4real/quantgen/blob/" + repo_ref + "/" + target[6:]
        elif target.startswith("../") and ".." not in target[3:]:
            url = "https://github.com/xkg4real/quantgen/blob/" + repo_ref + "/docs/" + target[3:]
        else:
            return label + " (" + target + ")"
        return '<a href="' + url + '">' + label + "</a>"
    value = re.sub(r"\[([^\]]+)\]\(([^\s)]+)\)", link, value)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    return re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)


def render_markdown(text: str, repo_ref: str) -> str:
    lines = text.splitlines()
    output, i = [], 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            output.append("<pre>" + html.escape("\n".join(block)) + "</pre>")
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            output.append('<div class="table-wrap"><table>')
            for row_no, row in enumerate(rows):
                tag = "th" if row_no == 0 else "td"
                output.append("<tr>" + "".join("<" + tag + ">" + inline(c, repo_ref) + "</" + tag + ">" for c in row) + "</tr>")
            output.append("</table></div>")
            continue
        elif line.startswith(">"):
            quoted = []
            while i < len(lines) and lines[i].startswith(">"):
                quoted.append(inline(lines[i].removeprefix(">").strip(), repo_ref))
                i += 1
            output.append("<blockquote>" + "<br>".join(quoted) + "</blockquote>")
            continue
        elif re.match(r"^#{1,6} ", line):
            level = min(len(line.split(" ", 1)[0]) + 1, 6)
            output.append(f"<h{level}>" + inline(line.split(" ", 1)[1], repo_ref) + f"</h{level}>")
        elif line.startswith(("- ", "* ")) or re.match(r"^\d+\. ", line):
            items = []
            while i < len(lines) and (lines[i].startswith(("- ", "* ")) or re.match(r"^\d+\. ", lines[i])):
                items.append(re.sub(r"^(?:[-*] |\d+\. )", "", lines[i]))
                i += 1
            output.append("<ul>" + "".join("<li>" + inline(x, repo_ref) + "</li>" for x in items) + "</ul>")
            continue
        elif line.strip():
            paragraph = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(?:#|>|\||```|- |\* |\d+\. )", lines[i]):
                paragraph.append(lines[i])
                i += 1
            output.append("<p>" + inline(" ".join(paragraph), repo_ref) + "</p>")
            continue
        i += 1
    return "\n".join(output)


def write_word(text: str, output: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(.7)
        section.bottom_margin = Inches(.7)
        section.left_margin = Inches(.65)
        section.right_margin = Inches(.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Malgun Gothic"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.styles["Title"].font.size = Pt(20)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.add_run("QUANTGEN · 검토용 초안 · ")
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        footer._p.append(page)
        for run in footer.runs:
            run.font.size = Pt(8)

    def plain(s):
        s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", s)
        return s.replace("**", "").replace("`", "")
    lines, i = text.splitlines(), 0
    in_code = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            in_code = not in_code
        elif in_code:
            doc.add_paragraph(line)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", x) for x in row):
                    rows.append(row)
                i += 1
            width = max(map(len, rows))
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            for row_id, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, content in zip(cells, row):
                    cell.text = plain(content)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                            run.bold = row_id == 0
                if row_id == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
            continue
        elif re.match(r"^#{1,6} ", line):
            level = len(line.split(" ", 1)[0])
            title = plain(line.split(" ", 1)[1])
            doc.add_heading(title, 0 if level == 1 else min(level - 1, 3))
        elif line.startswith(">"):
            doc.add_paragraph(plain(line.removeprefix(">").strip()))
        elif line.startswith("- "):
            doc.add_paragraph(plain(line[2:]), style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(plain(re.sub(r"^\d+\. ", "", line)), style="List Number")
        elif line.strip():
            doc.add_paragraph(plain(line))
        i += 1
    doc.core_properties.title = lines[0].removeprefix("# ")
    doc.core_properties.subject = "QUANTGEN 운영 정책 및 양식 · 검토용 초안"
    doc.core_properties.author = "QUANTGEN"
    doc.save(output)


def export(destination: Path, repo_ref: str) -> dict:
    current = collect(ROOT)
    if current["errors"]:
        raise ValueError("Fix document structure before exporting")
    report = check(ROOT)
    if report["consistency_status"] != "PASS":
        raise ValueError("Review current document/source changes before exporting")
    rows = registry(ROOT)
    # Validate all fixed inputs before creating an output directory.
    inputs = [row["path"] for row in rows]
    inputs += ["docs/compliance/" + name for name in SUPPLEMENTS]
    inputs += list(PIPELINE_FILES)
    for name in inputs:
        if not safe_path(ROOT, name).is_file():
            raise ValueError("Missing deliverable input: " + name)
    destination = destination.resolve()
    if destination.is_relative_to(ROOT) or destination.exists() or destination.is_symlink():
        raise ValueError("Use a new output directory outside the checkout")
    destination.mkdir(parents=True, exist_ok=False)
    source_dir = destination / "source"
    documents_dir, forms_dir = destination / "documents", destination / "forms"
    for directory in (source_dir, documents_dir, forms_dir):
        directory.mkdir()
    verification = destination / "verification"
    verification.mkdir()
    (verification / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (verification / "report.md").write_text(report_markdown(report), encoding="utf-8")
    for name in inputs:
        target = source_dir / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(safe_path(ROOT, name), target)
    toc, articles, field_rows = [], [], []
    for row in rows:
        path = safe_path(ROOT, row["path"])
        text = path.read_text(encoding="utf-8")
        form = form_section(text)
        name = path.stem
        write_word(text, documents_dir / (name + ".docx"))
        form_text = "# " + row["title"] + " 양식\n\n검토용 서식입니다. 실제 개인정보·증빙은 보호된 운영 기록에서 관리하고 저장소에 커밋하지 마세요.\n\n" + form + "\n"
        (forms_dir / (name + "-form.md")).write_text(form_text, encoding="utf-8")
        write_word(form_text, forms_dir / (name + "-form.docx"))
        toc.append(f'<li><a href="#doc-{row["id"]:02d}">{row["id"]:02d} {html.escape(row["title"])}</a></li>')
        articles.append(f'<article id="doc-{row["id"]:02d}">' + render_markdown(text, repo_ref) + "</article>")
        for field in sorted(set(re.findall(r"\{\{([^{}]+)\}\}", text))):
            field_rows.append([row["id"], row["title"], field, "", "공개 정책값과 개별 요청서 값은 별도 관리"])
    for name in SUPPLEMENTS:
        if name.endswith(".md"):
            key = name.removesuffix(".md").lower()
            toc.append(f'<li><a href="#{key}">{html.escape(name)}</a></li>')
            articles.append(f'<article id="{key}">' + render_markdown((ROOT / "docs/compliance" / name).read_text(encoding="utf-8"), repo_ref) + "</article>")
    toc.append('<li><a href="#verification">자동 점검 결과</a></li>')
    articles.append('<article id="verification">' + render_markdown(report_markdown(report), repo_ref) + '</article>')
    with (destination / "FIELD_INDEX.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["문서번호", "문서명", "치환필드", "확정값", "관리주의"])
        writer.writerows(field_rows)
    content = '''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>QUANTGEN 정책 20종과 양식</title><style>
body{font-family:system-ui,'Malgun Gothic',sans-serif;background:#f7f5f0;color:#202823;margin:0;line-height:1.75}header,main{max-width:1120px;margin:auto;padding:24px}h1{font-size:28px}h2{font-size:26px}h3{font-size:21px}h4{font-size:18px}a{color:#194c40}nav{background:#fffefb;padding:20px;border:1px solid #646e68}nav ul{columns:2}article{background:#fffefb;margin:24px 0;padding:30px;border:1px solid #646e68;scroll-margin:20px}code{background:#e5eee8;padding:2px 4px;overflow-wrap:anywhere}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f1f2eb;padding:15px}table{border-collapse:collapse;width:100%;font-size:14px}th,td{border:1px solid #646e68;padding:9px;text-align:left;vertical-align:top;overflow-wrap:anywhere}th{background:#e5eee8}.table-wrap{overflow-x:auto}input,button{font:inherit;padding:10px;border:2px solid #194c40;background:white;color:#202823}input{width:min(600px,80%)}:focus-visible{outline:3px solid #194c40;outline-offset:3px}.status{background:#fff3cf;padding:16px;border:1px solid #725100}#count{font-size:14px}@media(max-width:700px){nav ul{columns:1}article{padding:16px}}@media print{body{background:white}header input,header button,nav,#count{display:none}article{break-before:page;border:0;margin:0;padding:0}article[hidden]{display:block}a{color:#202823}table{font-size:9pt}h2,h3,h4{break-after:avoid}tr{break-inside:avoid}}
</style><header><h1>QUANTGEN 운영 정책 20종과 양식</h1><p class="status">2026-10-31 한국시간 기준 · 검토용 초안 · 공개 승인 전. 사업자 정보·운영 증빙·미구현 항목은 문서와 RELEASE_GAPS에서 확인하세요.</p><p>Word 문서는 documents, 별도 작성 양식은 forms, 저장소용 원본과 파이프라인은 source 폴더에 있습니다. 이 묶음은 QUANTGEN 체크아웃에 적용하는 변경 묶음이며 프로그램 전체를 포함하지 않습니다. 검색과 인쇄는 네트워크 없이 동작합니다.</p><label for="q">문서 내용 검색 </label><input id="q" type="search" autocomplete="off"><button onclick="window.print()">전체 인쇄</button><p id="count" role="status"></p><nav aria-label="문서 목록"><ul>'''
    content += "".join(toc) + "</ul></nav></header><main>" + "".join(articles)
    content += '''</main><script>const q=document.getElementById('q');q.addEventListener('input',()=>{let n=0;for(const a of document.querySelectorAll('article')){a.hidden=!a.textContent.toLowerCase().includes(q.value.toLowerCase());if(!a.hidden)n++;}document.getElementById('count').textContent=n+'개 문서 표시';});</script></html>'''
    (destination / "QUANTGEN_20_DOCUMENTS.html").write_text(content, encoding="utf-8")
    manifest = {"schema_version": 1, "document_count": len(rows), "word_document_count": len(rows),
                "separate_form_count": len(rows), "source_sha256": {row["path"]: current["documents"][str(row["id"])]["sha256"] for row in rows},
                "scope": "review drafts and reusable forms; not publication/legal approval", "files": {}}
    for file in sorted(destination.rglob("*")):
        if file.is_file():
            manifest["files"][file.relative_to(destination).as_posix()] = hashlib.sha256(file.read_bytes()).hexdigest()
    (destination / "MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    archive_path = destination / "QUANTGEN_Compliance_20_Documents.zip"
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for file in sorted(destination.rglob("*")):
            if file.is_file() and file != archive_path:
                archive.write(file, file.relative_to(destination))
    return {"documents": len(rows), "separate_forms": len(rows), "word_files": 2 * len(rows), "archive": str(archive_path)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--repository-ref", default="gpt/compliance-docs-20261031")
    args = parser.parse_args(argv)
    if not re.fullmatch(r"[A-Za-z0-9_./-]+", args.repository_ref) or ".." in args.repository_ref:
        parser.error("Invalid repository ref")
    print(json.dumps(export(args.output, args.repository_ref), ensure_ascii=False))


if __name__ == "__main__":
    main()
